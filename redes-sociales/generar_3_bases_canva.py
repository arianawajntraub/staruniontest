#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
import subprocess
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = "/Users/arianawajntraub/Documents/antigravity/star-union"
OUT_DIR = os.path.join(BASE_DIR, "redes-sociales")
PIEZAS_DIR = os.path.join(OUT_DIR, "piezas")
BASES_DIR = os.path.join(PIEZAS_DIR, "3_BASES_MAESTRAS_CANVA")
TMP_DIR = os.path.join(OUT_DIR, "tmp_render")
DRIVE_DIR = "/Users/arianawajntraub/Library/CloudStorage/GoogleDrive-ariana@starunionservicios.com.ar/Mi unidad/su para AG/REDES_SOCIALES_Y_SHORTS_2026"
DRIVE_BASES_DIR = os.path.join(DRIVE_DIR, "3_BASES_MAESTRAS_CANVA")

os.makedirs(BASES_DIR, exist_ok=True)
os.makedirs(DRIVE_BASES_DIR, exist_ok=True)
os.makedirs(TMP_DIR, exist_ok=True)

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# Convert px at 96 DPI to Inches for PPTX (1px = 1/96 inch)
def px(val):
    return Inches(val / 96.0)

def pt_px(val):
    # In 96dpi slide, 1px = 0.75 pt
    return Pt(val * 0.75)

GREEN = RGBColor(3, 88, 45)       # #03582D
DARK_GREEN = RGBColor(19, 53, 34) # #133522
YELLOW = RGBColor(255, 204, 41)   # #FFCC29
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(17, 24, 39)        # #111827
MUTED = RGBColor(55, 65, 81)      # #374151
LIGHT_BG = RGBColor(243, 244, 246)# #F3F4F6
BORDER_GRAY = RGBColor(229, 231, 235)
SAFE_GUIDE = RGBColor(220, 38, 38)

from generar_pack_redes import POSTS

def build_native_pptx(width_px, height_px, safe_top, safe_bottom, safe_left, safe_right, mode_label, out_pptx):
    """
    Creates a 5-slide native Editable PowerPoint (.pptx) where each slide is one of the 5 Star Union posts,
    respecting the exact Safe Zones (Márgenes de Seguridad) of the network.
    When uploaded to Canva, every text box, photo, logo, and pill is 100% natively editable in Canva,
    and Ariana can add the brand music once and duplicate slides!
    """
    prs = Presentation()
    prs.slide_width = px(width_px)
    prs.slide_height = px(height_px)
    blank_layout = prs.slide_layouts[6]

    logo_path = os.path.join(BASE_DIR, "assets/img/logo-star-union.png")
    tuv_path = os.path.join(BASE_DIR, "assets/img/tuv-iso-9001-badge.png")

    safe_w = width_px - safe_left - safe_right
    safe_h = height_px - safe_top - safe_bottom

    for post in POSTS:
        slide = prs.slides.add_slide(blank_layout)

        # 1. White Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(0), px(0), px(width_px), px(height_px))
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE
        bg.line.fill.background()

        # 2. Top Green Accent Strip inside Safe Zone top
        top_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(safe_left), px(safe_top), px(safe_w), px(10))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = GREEN
        top_bar.line.fill.background()

        # 3. Logo Star Union (Left inside Safe Zone)
        cur_y = safe_top + 24
        slide.shapes.add_picture(logo_path, px(safe_left), px(cur_y), height=px(58))

        # 4. ISO 9001 Box (Right inside Safe Zone)
        iso_w = 290
        iso_x = safe_left + safe_w - iso_w
        iso_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(iso_x), px(cur_y - 4), px(iso_w), px(64))
        iso_box.fill.solid()
        iso_box.fill.fore_color.rgb = LIGHT_BG
        iso_box.line.color.rgb = BORDER_GRAY
        slide.shapes.add_picture(tuv_path, px(iso_x + 12), px(cur_y + 6), height=px(42))

        iso_tb = slide.shapes.add_textbox(px(iso_x + 68), px(cur_y + 4), px(iso_w - 74), px(54))
        tf_iso = iso_tb.text_frame
        tf_iso.word_wrap = True
        p1 = tf_iso.paragraphs[0]
        p1.text = "ISO 9001:2015"
        p1.font.name = "Open Sans"
        p1.font.size = pt_px(14)
        p1.font.bold = True
        p1.font.color.rgb = GREEN
        p2 = tf_iso.add_paragraph()
        p2.text = "TÜV Rheinland · ID 9000041847"
        p2.font.name = "Open Sans"
        p2.font.size = pt_px(12)
        p2.font.color.rgb = MUTED

        # 5. Category Pill
        cur_y += 74
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(safe_left), px(cur_y), px(420), px(36))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(230, 239, 234)
        pill.line.fill.background()
        tf_pill = pill.text_frame
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = post["tag"]
        p_pill.font.name = "Open Sans"
        p_pill.font.size = pt_px(14)
        p_pill.font.bold = True
        p_pill.font.color.rgb = GREEN

        # 6. Editable Main Title
        cur_y += 46
        title_h = 115 if height_px <= 1350 else 130
        title_tb = slide.shapes.add_textbox(px(safe_left), px(cur_y), px(safe_w), px(title_h))
        tf_t = title_tb.text_frame
        tf_t.word_wrap = True
        pt_t = tf_t.paragraphs[0]
        pt_t.text = post["title"]
        pt_t.font.name = "Open Sans"
        pt_t.font.size = pt_px(36 if height_px <= 1350 else 40)
        pt_t.font.bold = True
        pt_t.font.color.rgb = INK

        # 7. Editable Subtitle
        cur_y += title_h + 4
        sub_h = 72 if height_px <= 1350 else 86
        sub_tb = slide.shapes.add_textbox(px(safe_left), px(cur_y), px(safe_w), px(sub_h))
        tf_s = sub_tb.text_frame
        tf_s.word_wrap = True
        pt_s = tf_s.paragraphs[0]
        pt_s.text = post["subtitle"]
        pt_s.font.name = "Open Sans"
        pt_s.font.size = pt_px(20 if height_px <= 1350 else 22)
        pt_s.font.color.rgb = MUTED

        # 8. KPI Pills Row
        cur_y += sub_h + 10
        kpi_w = int((safe_w - 24) / 3)
        for idx, kpi in enumerate(post["kpis"]):
            kx = safe_left + idx * (kpi_w + 12)
            k_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(kx), px(cur_y), px(kpi_w), px(38))
            k_shape.fill.solid()
            k_shape.fill.fore_color.rgb = LIGHT_BG
            k_shape.line.color.rgb = BORDER_GRAY
            ktf = k_shape.text_frame
            kp = ktf.paragraphs[0]
            kp.text = f"✓ {kpi}"
            kp.font.name = "Open Sans"
            kp.font.size = pt_px(14)
            kp.font.bold = True
            kp.font.color.rgb = DARK_GREEN

        # 9. Photo Frame (Sized to fit strictly inside Safe Zone above Footer)
        cur_y += 52
        footer_h = 74
        footer_y = safe_top + safe_h - footer_h
        avail_photo_h = max(260, footer_y - cur_y - 18)

        photo_abs = os.path.normpath(os.path.join(PIEZAS_DIR, post["photo"]))
        # Crop photo to exact aspect ratio of (safe_w x avail_photo_h) so PPTX/Canva doesn't stretch it
        cropped_photo_path = os.path.join(TMP_DIR, f"crop_{width_px}x{height_px}_{post['id']}.jpg")
        with Image.open(photo_abs) as im:
            im = im.convert("RGB")
            iw, ih = im.size
            target_ratio = safe_w / float(avail_photo_h)
            cur_ratio = iw / float(ih)
            if cur_ratio > target_ratio:
                new_w = int(ih * target_ratio)
                left = int((iw - new_w) / 2)
                im_cropped = im.crop((left, 0, left + new_w, ih))
            else:
                new_h = int(iw / target_ratio)
                pos_pct = 0.43 if post["id"] in ["01", "03"] else 0.25
                top = int((ih - new_h) * pos_pct)
                im_cropped = im.crop((0, top, iw, top + new_h))
            im_cropped.save(cropped_photo_path, "JPEG", quality=95)

        slide.shapes.add_picture(cropped_photo_path, px(safe_left), px(cur_y), width=px(safe_w), height=px(avail_photo_h))

        # 10. Footer inside Safe Zone
        foot_tb = slide.shapes.add_textbox(px(safe_left), px(footer_y + 8), px(safe_w - 300), px(62))
        ftf = foot_tb.text_frame
        ftf.word_wrap = True
        fp1 = ftf.paragraphs[0]
        fp1.text = "Star Union S.R.L. · Base Guaminí 5657, CABA"
        fp1.font.name = "Open Sans"
        fp1.font.size = pt_px(16)
        fp1.font.bold = True
        fp1.font.color.rgb = INK
        fp2 = ftf.add_paragraph()
        fp2.text = "WhatsApp B2B: +54 9 11 2275-9369 · 0810-333-8646"
        fp2.font.name = "Open Sans"
        fp2.font.size = pt_px(14)
        fp2.font.color.rgb = MUTED

        cta_w = 285
        cta_btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(safe_left + safe_w - cta_w), px(footer_y + 12), px(cta_w), px(48))
        cta_btn.fill.solid()
        cta_btn.fill.fore_color.rgb = GREEN
        cta_btn.line.color.rgb = YELLOW
        cta_btn.line.width = Pt(2)
        ctf = cta_btn.text_frame
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cp.text = "www.starunion.com.ar"
        cp.font.name = "Open Sans"
        cp.font.size = pt_px(16)
        cp.font.bold = True
        cp.font.color.rgb = WHITE

    prs.save(out_pptx)


def render_safe_html_and_files():
    """
    Also updates the HTML-rendered JPGs & Editable PDFs so that:
    1) BASE 1 REEL (1080x1920): Safe Zone Top 250px, Bottom 420px, Left 84px, Right 140px (for Reels/Shorts right-side buttons!)
    2) BASE 2 FACEBOOK / IG FEED (1080x1350): Safe Zone Top 145px, Bottom 145px, Left 84px, Right 84px (so 1:1 profile grid never cuts logo or footer!)
    3) BASE 3 LINKEDIN (1200x1200): Safe Zone 90px all around.
    """
    configs = [
        {
            "name": "BASE_1_REEL_SHORT_1080x1920",
            "w": 1080, "h": 1920,
            "st": 250, "sb": 420, "sl": 84, "sr": 140,
            "label": "REEL / SHORT / STORY (1080×1920 · Margen Seguro: 250px arriba, 420px abajo, 140px derecha)"
        },
        {
            "name": "BASE_2_FACEBOOK_IG_FEED_1080x1350",
            "w": 1080, "h": 1350,
            "st": 145, "sb": 145, "sl": 84, "sr": 84,
            "label": "FACEBOOK / IG FEED (1080×1350 · Margen Seguro: 145px arriba/abajo para que la grilla 1:1 no corte nada)"
        },
        {
            "name": "BASE_3_LINKEDIN_1200x1200",
            "w": 1200, "h": 1200,
            "st": 90, "sb": 90, "sl": 90, "sr": 90,
            "label": "LINKEDIN B2B (1200×1200 · Margen Seguro Corporativo 90px en los 4 lados)"
        }
    ]

    for cfg in configs:
        out_pptx = os.path.join(BASES_DIR, f"{cfg['name']}_EDITABLE_CANVA.pptx")
        build_native_pptx(cfg["w"], cfg["h"], cfg["st"], cfg["sb"], cfg["sl"], cfg["sr"], cfg["label"], out_pptx)

    # Also render Post #1 in all 3 bases as JPG (with and without Safe Zone overlay guide) + Vector PDF for Canva
    post = POSTS[0]
    for cfg in configs:
        w, h = cfg["w"], cfg["h"]
        st, sb, sl, sr = cfg["st"], cfg["sb"], cfg["sl"], cfg["sr"]
        title_fs = 40 if h >= 1350 else 38
        sub_fs = 21 if h >= 1350 else 20
        photo_h = 460 if h == 1920 else (380 if h == 1350 else 340)

        kpis_html = "".join([
            f'<span style="background:#f3f4f6;border:1px solid #e5e7eb;color:#133522;font-weight:800;font-size:15px;padding:8px 14px;border-radius:999px;">✓ {k}</span>'
            for k in post["kpis"]
        ])

        for show_guide in [False, True]:
            guide_html = ""
            if show_guide:
                guide_html = f"""
                <div style="position:fixed;inset:0;pointer-events:none;z-index:99;">
                  <div style="position:absolute;top:0;left:0;right:0;height:{st}px;background:rgba(220,38,38,0.10);border-bottom:3px dashed #dc2626;display:flex;align-items:center;justify-content:center;font-size:18px;font-weight:800;color:#dc2626;">
                    ZONA PROTEGIDA SUPERIOR ({st}px · Interfaz / Recorte Grilla 1:1)
                  </div>
                  <div style="position:absolute;bottom:0;left:0;right:0;height:{sb}px;background:rgba(220,38,38,0.10);border-top:3px dashed #dc2626;display:flex;align-items:center;justify-content:center;font-size:18px;font-weight:800;color:#dc2626;">
                    ZONA PROTEGIDA INFERIOR ({sb}px · Copy, Audio/Música y Recorte Grilla)
                  </div>
                  <div style="position:absolute;top:{st}px;bottom:{sb}px;right:0;width:{sr}px;background:rgba(220,38,38,0.08);border-left:3px dashed #dc2626;display:flex;align-items:center;justify-content:center;writing-mode:vertical-rl;font-size:16px;font-weight:800;color:#dc2626;">
                    MARGEN DERECHO ({sr}px · Botonera Reels/Shorts)
                  </div>
                  <div style="position:absolute;top:{st}px;bottom:{sb}px;left:0;width:{sl}px;background:rgba(220,38,38,0.08);border-right:3px dashed #dc2626;"></div>
                </div>
                """

            html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<link href="https://fonts.googleapis.com/css2?family=Open+Sans:wght@400;600;700;800&display=swap" rel="stylesheet">
<style>
  @page {{ size: {w}px {h}px; margin: 0; }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    width: {w}px;
    height: {h}px;
    overflow: hidden;
    background: #ffffff;
    font-family: 'Open Sans', sans-serif;
    color: #111827;
    padding: {st}px {sr}px {sb}px {sl}px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
  }}
</style>
</head>
<body>
  {guide_html}
  <div style="border-top:10px solid #03582d;padding-top:20px;">
    <div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:18px;">
      <img src="../../assets/img/logo-star-union.png" style="height:58px;width:auto;object-fit:contain;" alt="Star Union">
      <div style="display:flex;align-items:center;gap:10px;background:#f9fafb;border:1px solid #e5e7eb;padding:8px 14px;border-radius:12px;">
        <img src="../../assets/img/tuv-iso-9001-badge.png" style="height:36px;width:auto;object-fit:contain;" alt="ISO 9001">
        <div>
          <div style="font-size:12px;font-weight:800;color:#03582d;">ISO 9001:2015</div>
          <div style="font-size:11px;color:#4b5563;font-weight:700;">TÜV Rheinland · ID 9000041847</div>
        </div>
      </div>
    </div>

    <div style="display:inline-block;background:rgba(3,88,45,0.09);color:#03582d;font-weight:800;font-size:14px;letter-spacing:0.07em;padding:6px 14px;border-radius:999px;margin-bottom:12px;">
      {post["tag"]}
    </div>

    <h1 style="font-size:{title_fs}px;line-height:1.13;font-weight:800;color:#111827;letter-spacing:-0.02em;margin-bottom:10px;">
      {post["title"]}
    </h1>

    <p style="font-size:{sub_fs}px;line-height:1.38;color:#374151;font-weight:600;margin-bottom:14px;">
      {post["subtitle"]}
    </p>

    <div style="display:flex;gap:8px;flex-wrap:wrap;">
      {kpis_html}
    </div>
  </div>

  <div style="width:100%;height:{photo_h}px;border-radius:22px;overflow:hidden;border:2px solid #e5e7eb;background:#f9fafb;">
    <img src="{post['photo']}" style="width:100%;height:100%;object-fit:cover;object-position:center 43%;" alt="">
  </div>

  <div style="display:flex;align-items:center;justify-content:space-between;padding-top:12px;border-top:2px solid #f3f4f6;">
    <div>
      <div style="font-size:16px;font-weight:800;color:#111827;">Star Union S.R.L. · Base Guaminí 5657, CABA</div>
      <div style="font-size:14px;font-weight:600;color:#4b5563;">WhatsApp B2B: +54 9 11 2275-9369 · 0810-333-8646</div>
    </div>
    <div style="background:#03582d;color:#ffffff;font-weight:800;font-size:16px;padding:11px 22px;border-radius:999px;border-bottom:3px solid #ffcc29;">
      www.starunion.com.ar
    </div>
  </div>
</body>
</html>"""
            suffix = "_CON_GUIA_MARGENES" if show_guide else ""
            html_path = os.path.join(TMP_DIR, f"{cfg['name']}{suffix}.html")
            png_path = os.path.join(TMP_DIR, f"{cfg['name']}{suffix}.png")
            jpg_path = os.path.join(BASES_DIR, f"{cfg['name']}{suffix}.jpg")
            with open(html_path, "w", encoding="utf-8") as f:
                f.write(html)
            subprocess.run([
                CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                f"--window-size={w},{h}", f"--screenshot={png_path}", f"file://{html_path}"
            ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            Image.open(png_path).convert("RGB").save(jpg_path, "JPEG", quality=95)

            if not show_guide:
                pdf_path = os.path.join(BASES_DIR, f"{cfg['name']}_EDITABLE_CANVA.pdf")
                subprocess.run([
                    CHROME, "--headless=new", "--disable-gpu",
                    "--print-to-pdf-no-header", f"--print-to-pdf={pdf_path}", f"file://{html_path}"
                ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    # Copy all bases to local bases-canva-editables and both Google Drive folders
    local_bases_edit = os.path.join(OUT_DIR, "bases-canva-editables")
    drive_madre_dir = os.path.join(DRIVE_DIR, "00_BASES_MADRE_CANVA_CON_MARGEN_SEGURO")
    os.makedirs(local_bases_edit, exist_ok=True)
    os.makedirs(drive_madre_dir, exist_ok=True)
    subprocess.run(f'cp -R "{BASES_DIR}/"* "{DRIVE_BASES_DIR}/"', shell=True, check=True)
    subprocess.run(f'cp -R "{BASES_DIR}/"*.pptx "{local_bases_edit}/"', shell=True, check=True)
    subprocess.run(f'cp -R "{BASES_DIR}/"*.pptx "{drive_madre_dir}/"', shell=True, check=True)
    print("3 Master Bases (PPTX + PDF + Safe Zone JPGs) created and synced to Google Drive.")

if __name__ == "__main__":
    render_safe_html_and_files()
